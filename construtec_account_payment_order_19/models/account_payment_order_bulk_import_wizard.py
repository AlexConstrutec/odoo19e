import datetime
import re
import unicodedata

from odoo import _, api, fields, models
from odoo.exceptions import UserError

# Alias reconocidos por columna canónica - el mapeo automático compara el encabezado real del
# Excel (normalizado: minúsculas, sin acentos, espacios colapsados) contra esta lista; si no
# encuentra nada razonable, el usuario corrige a mano en el paso "Mapeo de Columnas" - el Excel
# real nunca viene con el mismo formato dos veces, así que esto es solo un punto de partida.
COLUMNAS_COMUNES_ALIAS = {
    'col_grupo': ['grupo', 'no de solicitud', 'no. de solicitud', 'solicitud', 'lote'],
    'col_solicitante': ['solicitante', 'jefe', 'jefe de tecnicos', 'jefe de equipo'],
    'col_cuenta_analitica': ['cuenta analitica', 'analitica', 'cuenta', 'proyecto', 'ubicacion'],
    'col_fecha': ['fecha'],
}
COLUMNAS_VIATICOS_ALIAS = {
    'col_periodo_del': ['periodo del', 'del', 'fecha inicio', 'inicio'],
    'col_periodo_al': ['periodo al', 'al', 'fecha fin', 'fin'],
    'col_depositar_directo': ['depositar directo', 'depositar directo a tecnicos', 'deposito directo'],
    'col_tecnico': ['tecnico', 'empleado', 'nombre del tecnico', 'nombre'],
    'col_cantidad': ['cantidad', 'dias', 'cantidad de dias'],
    'col_costo_individual': ['costo individual', 'costo', 'monto', 'monto diario'],
    'col_cuenta_acreditar': ['cuenta a acreditar', 'cuenta bancaria', 'no cuenta', 'numero de cuenta'],
    'col_banco': ['banco'],
    'col_tipo_cuenta': ['tipo de cuenta', 'tipo cuenta'],
}
COLUMNAS_MATERIALES_ALIAS = {
    'col_proveedor': ['proveedor'],
    'col_material': ['material', 'descripcion', 'producto'],
    'col_unidad': ['unidad', 'unidad de medida', 'uom'],
    'col_cantidad_material': ['cantidad'],
    'col_precio_estimado': ['precio estimado', 'precio', 'precio unitario'],
}


def _normalizar_texto(texto):
    if not texto:
        return ''
    texto = str(texto).strip().lower()
    texto = ''.join(
        c for c in unicodedata.normalize('NFD', texto) if unicodedata.category(c) != 'Mn')
    return re.sub(r'\s+', ' ', texto)


def _auto_mapear_columnas(encabezados_reales, alias_por_campo):
    """Adivina, para cada campo canónico, cuál encabezado real del Excel le corresponde -
    comparación por substring en ambos sentidos sobre texto normalizado. Nunca reutiliza el
    mismo encabezado real para dos campos distintos. Es solo un punto de partida editable -
    ver COLUMNAS_*_ALIAS más arriba."""
    resultado = {}
    usados = set()
    normalizados = [(h, _normalizar_texto(h)) for h in encabezados_reales]
    for campo, alias in alias_por_campo.items():
        for header, header_norm in normalizados:
            if header in usados or not header_norm:
                continue
            if any(a in header_norm or header_norm in a for a in alias):
                resultado[campo] = header
                usados.add(header)
                break
    return resultado


class AccountPaymentOrderBulkImportWizard(models.TransientModel):
    _name = 'account.payment.order.bulk.import.wizard'
    _description = 'Importar Órdenes de Pago desde Excel (Viáticos o Materiales)'

    state = fields.Selection([
        ('upload', 'Subir Archivo'),
        ('mapping', 'Mapeo de Columnas'),
        ('review', 'Revisión'),
    ], default='upload', required=True)
    tipo = fields.Selection([
        ('anticipo_viaticos', 'Solicitud de Viáticos'),
        ('anticipo_materiales', 'Solicitud de Materiales'),
    ], string='Tipo de Solicitud', required=True, default='anticipo_viaticos')
    excel_file = fields.Binary(string='Archivo Excel')
    excel_filename = fields.Char(string='Nombre del Archivo')
    detected_headers = fields.Char(string='Columnas Detectadas en el Archivo', readonly=True)

    # Mapeo - columnas comunes a ambos tipos.
    col_grupo = fields.Char(
        string='Columna: No. de Solicitud / Grupo',
        help='Filas con el mismo valor en esta columna arman UNA sola Orden de Pago (varios '
             'técnicos o varios materiales). Si se deja vacía, cada fila se vuelve su propia '
             'Orden independiente.')
    col_solicitante = fields.Char(
        string='Columna: Solicitante',
        help='Nombre del Jefe de Técnicos que pide la Orden - se busca contra los Contactos '
             'marcados como empleado.')
    col_cuenta_analitica = fields.Char(string='Columna: Cuenta Analítica')
    col_fecha = fields.Char(string='Columna: Fecha')

    # Mapeo - solo Viáticos.
    col_periodo_del = fields.Char(string='Columna: Período Del')
    col_periodo_al = fields.Char(string='Columna: Período Al')
    col_depositar_directo = fields.Char(string='Columna: ¿Depositar Directo a Técnicos?')
    col_tecnico = fields.Char(string='Columna: Técnico')
    col_cantidad = fields.Char(string='Columna: Cantidad')
    col_costo_individual = fields.Char(string='Columna: Costo Individual')
    col_cuenta_acreditar = fields.Char(string='Columna: Cuenta a Acreditar (opcional)')
    col_banco = fields.Char(string='Columna: Banco (opcional)')
    col_tipo_cuenta = fields.Char(string='Columna: Tipo de Cuenta (opcional)')

    # Mapeo - solo Materiales.
    col_proveedor = fields.Char(string='Columna: Proveedor')
    col_material = fields.Char(string='Columna: Material / Descripción')
    col_unidad = fields.Char(string='Columna: Unidad de Medida')
    col_cantidad_material = fields.Char(string='Columna: Cantidad')
    col_precio_estimado = fields.Char(string='Columna: Precio Estimado')

    line_ids = fields.One2many(
        'account.payment.order.bulk.import.wizard.line', 'wizard_id', string='Filas')
    resumen = fields.Text(string='Resumen', readonly=True)

    @api.model
    def default_get(self, fields_list):
        """Bloquea incluso ABRIR el wizard del lado Procesador (Enterprise) - ahí las Órdenes
        llegan por sincronización, nunca se crean directamente. El mismo chequeo se repite en
        `action_aplicar()` (defensa en profundidad, por si se llama por API/automatización sin
        pasar por este `default_get()`). También exige "Administrador" (`account.group_account_
        manager`, mismo grupo que ya usa `_check_es_administrador_contable()` para Aplicar/
        Conciliar/Cancelar en `account_payment_order.py`) - pedido explícito del usuario: la
        importación masiva es una acción de Administrador, en ambas ediciones."""
        self.env['account.payment.order']._check_es_administrador_contable()
        if self.env.company.payment_order_role == 'procesador':
            raise UserError(_(
                'La importación masiva de Órdenes de Pago solo aplica del lado Solicitante '
                '(Community) - en Enterprise las Órdenes llegan por sincronización, nunca se '
                'crean directamente aquí.'))
        return super().default_get(fields_list)

    def _leer_workbook(self):
        import base64
        import io

        import openpyxl

        self.ensure_one()
        if not self.excel_file:
            raise UserError(_('Sube un archivo Excel antes de continuar.'))
        return openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(self.excel_file)), data_only=True)

    def _leer_encabezados(self, sheet):
        primera_fila = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return [str(c).strip() for c in primera_fila if c not in (None, '')]

    def action_leer_columnas(self):
        """Sube el archivo -> lee solo la fila de encabezados -> adivina el mapeo -> pasa a
        'mapping' para que el usuario lo confirme/corrija antes de leer los datos reales."""
        self.ensure_one()
        wb = self._leer_workbook()
        encabezados = self._leer_encabezados(wb.active)
        if not encabezados:
            raise UserError(_(
                'No se encontraron encabezados en la primera fila de la hoja "%s".',
                wb.active.title))
        self.detected_headers = ', '.join(encabezados)
        alias = dict(COLUMNAS_COMUNES_ALIAS)
        alias.update(
            COLUMNAS_VIATICOS_ALIAS if self.tipo == 'anticipo_viaticos'
            else COLUMNAS_MATERIALES_ALIAS)
        for campo, valor in _auto_mapear_columnas(encabezados, alias).items():
            self[campo] = valor
        self.state = 'mapping'

    def _columnas_requeridas(self):
        comunes = ['col_solicitante', 'col_cuenta_analitica']
        if self.tipo == 'anticipo_viaticos':
            return comunes + ['col_tecnico', 'col_cantidad', 'col_costo_individual']
        return comunes + ['col_material', 'col_cantidad_material', 'col_precio_estimado']

    def _resolver_contacto_empleado(self, texto):
        """Mismo criterio de 'único candidato o nada' ya usado en
        `_autosugerir_proveedor_materiales_id()` - exacto primero, luego parecido SOLO si hay
        un único resultado; nunca adivina entre varios para no vincular al técnico equivocado."""
        if not texto:
            return self.env['res.partner']
        texto = str(texto).strip()
        Partner = self.env['res.partner']
        domain_base = [('employee', '=', True)]
        exacto = Partner.search(domain_base + [('name', '=ilike', texto)], limit=2)
        if len(exacto) == 1:
            return exacto
        parecidos = Partner.search(domain_base + [('name', 'ilike', texto)], limit=2)
        return parecidos if len(parecidos) == 1 else self.env['res.partner']

    def _resolver_cuenta_analitica(self, texto):
        if not texto:
            return self.env['account.analytic.account']
        texto = str(texto).strip()
        Analytic = self.env['account.analytic.account']
        exacto = Analytic.search([('name', '=ilike', texto)], limit=2)
        if len(exacto) == 1:
            return exacto
        parecidos = Analytic.search([('name', 'ilike', texto)], limit=2)
        return parecidos if len(parecidos) == 1 else self.env['account.analytic.account']

    def _parse_fecha(self, valor):
        if not valor:
            return False
        if isinstance(valor, datetime.datetime):
            return valor.date()
        if isinstance(valor, datetime.date):
            return valor
        try:
            return fields.Date.to_date(str(valor).strip())
        except Exception:
            return False

    def _parse_bool(self, valor):
        if isinstance(valor, bool):
            return valor
        if not valor:
            return False
        return _normalizar_texto(valor) in ('si', 'sí', 'yes', 'true', '1', 'x')

    def _parse_tipo_cuenta(self, valor):
        texto = _normalizar_texto(valor)
        if 'ahorro' in texto:
            return 'ahorro'
        if 'monetari' in texto:
            return 'monetaria'
        return False

    def action_analizar(self):
        """Confirma el mapeo -> lee todas las filas de datos -> resuelve Solicitante/Cuenta
        Analítica/Técnico contra registros reales -> arma la vista previa editable en
        `line_ids`. Nunca crea nada real todavía - eso solo pasa en `action_aplicar()`."""
        self.ensure_one()
        faltantes = [c for c in self._columnas_requeridas() if not self[c]]
        if faltantes:
            etiquetas = [self._fields[c].string for c in faltantes]
            raise UserError(_(
                'Faltan columnas obligatorias por mapear: %s') % ', '.join(etiquetas))

        wb = self._leer_workbook()
        sheet = wb.active
        encabezados = self._leer_encabezados(sheet)
        indice = {h: i for i, h in enumerate(encabezados)}

        def valor(fila, campo_col):
            nombre_columna = self[campo_col]
            if not nombre_columna or nombre_columna not in indice:
                return None
            i = indice[nombre_columna]
            v = fila[i] if i < len(fila) else None
            return v.strip() if isinstance(v, str) else v

        nuevas = []
        for num_fila, fila in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if all(c in (None, '') for c in fila):
                continue
            solicitante_texto = valor(fila, 'col_solicitante')
            cuenta_texto = valor(fila, 'col_cuenta_analitica')
            solicitante = self._resolver_contacto_empleado(solicitante_texto)
            analytic = self._resolver_cuenta_analitica(cuenta_texto)
            grupo_valor = valor(fila, 'col_grupo')
            vals = {
                'fila_excel': num_fila,
                'grupo': str(grupo_valor) if grupo_valor not in (None, '') else '',
                'solicitante_texto': solicitante_texto,
                'solicitante_id': solicitante.id if solicitante else False,
                'cuenta_analitica_texto': cuenta_texto,
                'analytic_account_id': analytic.id if analytic else False,
                'fecha': self._parse_fecha(valor(fila, 'col_fecha')),
            }
            motivos = []
            if not solicitante:
                motivos.append(_('Solicitante no encontrado: "%s"') % (solicitante_texto or ''))
            if not analytic:
                motivos.append(_('Cuenta Analítica no encontrada: "%s"') % (cuenta_texto or ''))
            if self.tipo == 'anticipo_viaticos':
                tecnico_texto = valor(fila, 'col_tecnico')
                tecnico = self._resolver_contacto_empleado(tecnico_texto)
                if not tecnico:
                    motivos.append(_('Técnico no encontrado: "%s"') % (tecnico_texto or ''))
                vals.update({
                    'periodo_del': self._parse_fecha(valor(fila, 'col_periodo_del')),
                    'periodo_al': self._parse_fecha(valor(fila, 'col_periodo_al')),
                    'depositar_directo_tecnicos': self._parse_bool(
                        valor(fila, 'col_depositar_directo')),
                    'tecnico_texto': tecnico_texto,
                    'employee_partner_id': tecnico.id if tecnico else False,
                    'cantidad': int(valor(fila, 'col_cantidad') or 1),
                    'costo_individual': float(valor(fila, 'col_costo_individual') or 0.0),
                    'cuenta_acreditar': valor(fila, 'col_cuenta_acreditar') or False,
                    'banco': valor(fila, 'col_banco') or False,
                    'tipo_cuenta': self._parse_tipo_cuenta(valor(fila, 'col_tipo_cuenta')),
                })
            else:
                vals.update({
                    'proveedor_materiales_name': valor(fila, 'col_proveedor') or '',
                    'material_description': valor(fila, 'col_material') or '',
                    'uom_name': valor(fila, 'col_unidad') or '',
                    'qty': float(valor(fila, 'col_cantidad_material') or 1.0),
                    'estimated_price': float(valor(fila, 'col_precio_estimado') or 0.0),
                })
            vals['estado'] = 'revisar' if motivos else 'ok'
            vals['motivo_revisar'] = '; '.join(motivos)
            nuevas.append((0, 0, vals))

        if not nuevas:
            raise UserError(_('El archivo no tiene filas de datos para importar.'))

        self.line_ids = [(5, 0, 0)] + nuevas
        total = len(nuevas)
        con_revisar = len([n for n in nuevas if n[2]['estado'] == 'revisar'])
        self.resumen = _(
            '%(total)s fila(s) leídas: %(ok)s listas para crear, %(revisar)s necesitan '
            'revisión (Solicitante/Cuenta Analítica/Técnico sin resolver).',
            total=total, ok=total - con_revisar, revisar=con_revisar)
        self.state = 'review'

    def action_volver(self):
        self.ensure_one()
        self.line_ids = [(5, 0, 0)]
        self.state = 'mapping'

    def action_aplicar(self):
        """Crea las Órdenes de Pago reales, agrupando `line_ids` (solo las marcadas Incluir)
        por `grupo` - todas las filas del mismo grupo arman UNA Orden con varias líneas. Cada
        Orden queda en 'borrador' (nunca se auto-envía) para que el Solicitante la revise antes
        de Enviar - mismo criterio que "Cargar Cotización"."""
        self.ensure_one()
        self.env['account.payment.order']._check_es_administrador_contable()
        if self.env.company.payment_order_role == 'procesador':
            raise UserError(_(
                'La importación masiva de Órdenes de Pago solo aplica del lado Solicitante '
                '(Community) - en Enterprise las Órdenes llegan por sincronización, nunca se '
                'crean directamente aquí.'))
        lineas = self.line_ids.filtered('incluir')
        if not lineas:
            raise UserError(_('Selecciona al menos una fila (columna "Incluir") antes de aplicar.'))
        sin_resolver = lineas.filtered(lambda l: l.estado == 'revisar')
        if sin_resolver:
            raise UserError(_(
                'Hay %(cantidad)s fila(s) marcadas "Revisar" - corrige el Solicitante/Cuenta '
                'Analítica/Técnico antes de aplicar, o desmárcalas de "Incluir".',
                cantidad=len(sin_resolver)))

        Order = self.env['account.payment.order']
        grupos = {}
        for linea in lineas:
            clave = linea.grupo or f'__fila_{linea.id}'
            grupos.setdefault(clave, self.env['account.payment.order.bulk.import.wizard.line'])
            grupos[clave] |= linea

        ordenes_creadas = Order
        for lineas_grupo in grupos.values():
            primera = lineas_grupo[0]
            vals = {
                'tipo': self.tipo,
                'partner_id': primera.solicitante_id.id,
                'analytic_account_id': primera.analytic_account_id.id,
                'fecha': primera.fecha or fields.Date.context_today(self),
            }
            if self.tipo == 'anticipo_viaticos':
                vals.update({
                    'periodo_del': primera.periodo_del,
                    'periodo_al': primera.periodo_al,
                    'depositar_directo_tecnicos': primera.depositar_directo_tecnicos,
                    'viaticos_line_ids': [(0, 0, {
                        'employee_partner_id': l.employee_partner_id.id,
                        'cantidad': l.cantidad,
                        'costo_individual': l.costo_individual,
                        **({'cuenta_acreditar': l.cuenta_acreditar} if l.cuenta_acreditar else {}),
                        **({'banco': l.banco} if l.banco else {}),
                        **({'tipo_cuenta': l.tipo_cuenta} if l.tipo_cuenta else {}),
                    }) for l in lineas_grupo],
                })
            else:
                vals.update({
                    'proveedor_materiales_name': primera.proveedor_materiales_name or '',
                    'material_line_ids': [(0, 0, {
                        'product_name': l.material_description,
                        'description': l.material_description,
                        'uom_name': l.uom_name,
                        'qty': l.qty,
                        'estimated_price': l.estimated_price,
                        'vendor_name': primera.proveedor_materiales_name or '',
                    }) for l in lineas_grupo],
                })
            ordenes_creadas |= Order.create(vals)

        for orden in ordenes_creadas:
            orden.message_post(body=_('Creada por importación masiva desde Excel (%s).')
                                % (self.excel_filename or ''))
        return {
            'type': 'ir.actions.act_window',
            'name': _('Órdenes Creadas'),
            'res_model': 'account.payment.order',
            'view_mode': 'list,form',
            'domain': [('id', 'in', ordenes_creadas.ids)],
        }
