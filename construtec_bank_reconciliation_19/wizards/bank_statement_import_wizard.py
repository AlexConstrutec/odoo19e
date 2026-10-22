import base64
import csv
import io
import re
import unicodedata
from datetime import date, datetime, timedelta

from odoo import fields, models
from odoo.exceptions import UserError

# Palabras que identifican cada columna en el encabezado del estado de cuenta (sin tildes,
# minúsculas). El primer encabezado que contenga alguna de ellas gana.
COLUMN_KEYWORDS = {
    'date': ['fecha', 'date'],
    'ref': ['documento', 'referencia', 'no. doc', 'num doc', 'cheque', 'boleta', 'ref'],
    'name': ['descripcion', 'concepto', 'detalle', 'observacion', 'transaccion'],
    'debit': ['debito', 'cargo', 'retiro', 'egreso'],
    'credit': ['credito', 'abono', 'deposito', 'ingreso'],
    'amount': ['monto', 'importe', 'valor'],
    'balance': ['saldo', 'balance'],
}


def _norm(value):
    text = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', text).strip().lower()


def _parse_amount(value):
    if value in (None, ''):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    negative = text.startswith('(') and text.endswith(')') or text.startswith('-') \
        or text.endswith('-')
    text = re.sub(r'[^\d.,]', '', text)
    if not text:
        return 0.0
    if ',' in text and '.' in text:
        # El último separador es el decimal.
        if text.rfind(',') > text.rfind('.'):
            text = text.replace('.', '').replace(',', '.')
        else:
            text = text.replace(',', '')
    elif ',' in text:
        parts = text.split(',')
        text = text.replace(',', '.') if len(parts[-1]) in (1, 2) else text.replace(',', '')
    amount = float(text)
    return -amount if negative else amount


def _parse_date(value):
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        # Número de serie de Excel
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    text = str(value).strip().split(' ')[0]
    for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d/%m/%y', '%d-%m-%y', '%Y/%m/%d', '%d.%m.%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


class ConstrutecBankStatementImportWizard(models.TransientModel):
    _name = 'construtec.bank.statement.import.wizard'
    _description = 'Importar estado de cuenta (Construtec)'

    reconciliation_id = fields.Many2one('construtec.bank.reconciliation', required=True,
                                        readonly=True)
    file = fields.Binary(string='Archivo', required=True)
    filename = fields.Char()
    invert_sign = fields.Boolean(
        string='Invertir débito y crédito',
        help='Marque si el archivo viene desde el punto de vista contable (Cargo = depósito) '
             'en lugar del punto de vista del banco (Débito = retiro).')
    replace_unmatched = fields.Boolean(
        string='Reemplazar líneas sin empatar',
        help='Borra las líneas sin empatar de esta conciliación antes de importar.')
    only_period = fields.Boolean(string='Solo movimientos del periodo', default=True)

    def _read_rows(self):
        content = base64.b64decode(self.file)
        name = (self.filename or '').lower()
        if name.endswith('.xlsx'):
            import openpyxl
            book = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
            sheet = book.worksheets[0]
            return [list(row) for row in sheet.iter_rows(values_only=True)]
        if name.endswith('.xls'):
            import xlrd
            book = xlrd.open_workbook(file_contents=content)
            sheet = book.sheet_by_index(0)
            rows = []
            for r in range(sheet.nrows):
                row = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    else:
                        row.append(cell.value)
                rows.append(row)
            return rows
        if name.endswith('.csv') or name.endswith('.txt'):
            for encoding in ('utf-8-sig', 'latin-1'):
                try:
                    text = content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            # El sniffer de csv se confunde con montos como "1,000.00": se prueba cada separador y
            # se queda el primero con el que se encuentra el encabezado.
            for delimiter in (';', '\t', '|', ','):
                rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
                if self._find_header(rows):
                    return rows
            return list(csv.reader(io.StringIO(text)))
        if name.endswith('.pdf'):
            raise UserError('La lectura de PDF por banco todavía no está disponible. Por ahora '
                            'importe el Excel o CSV que descarga el banco; el PDF puede adjuntarse '
                            'en el campo "Estado de cuenta (PDF)".')
        raise UserError('Formato no reconocido. Use un archivo .xlsx, .xls o .csv.')

    def _detect_header(self, rows):
        found = self._find_header(rows)
        if not found:
            raise UserError('No se encontró el encabezado del estado de cuenta. Debe tener al '
                            'menos columnas de Fecha y Débito/Crédito (o Monto).')
        return found

    @staticmethod
    def _find_header(rows):
        for index, row in enumerate(rows[:40]):
            cells = [_norm(c) for c in row]
            mapping = {}
            for key, words in COLUMN_KEYWORDS.items():
                for col, cell in enumerate(cells):
                    if col in mapping.values() or not cell:
                        continue
                    if any(word in cell for word in words):
                        mapping[key] = col
                        break
            has_amounts = ('debit' in mapping and 'credit' in mapping) or 'amount' in mapping
            if 'date' in mapping and has_amounts:
                return index, mapping
        return None

    def action_import(self):
        self.ensure_one()
        rec = self.reconciliation_id
        if rec.state != 'draft':
            raise UserError('La conciliación ya está cerrada.')
        rows = self._read_rows()
        header_index, mapping = self._detect_header(rows)

        def cell(row, key):
            col = mapping.get(key)
            return row[col] if col is not None and col < len(row) else None

        vals_list, skipped = [], 0
        for row in rows[header_index + 1:]:
            line_date = _parse_date(cell(row, 'date'))
            if not line_date:
                continue
            if 'debit' in mapping and 'credit' in mapping:
                debit = abs(_parse_amount(cell(row, 'debit')))
                credit = abs(_parse_amount(cell(row, 'credit')))
            else:
                amount = _parse_amount(cell(row, 'amount'))
                debit, credit = (-amount, 0.0) if amount < 0 else (0.0, amount)
            if self.invert_sign:
                debit, credit = credit, debit
            if not debit and not credit:
                continue
            if self.only_period and not (rec.date_from <= line_date <= rec.date_to):
                skipped += 1
                continue
            ref = cell(row, 'ref')
            if isinstance(ref, float) and ref.is_integer():
                ref = int(ref)
            vals_list.append({
                'reconciliation_id': rec.id,
                'date': line_date,
                'ref': str(ref).strip() if ref not in (None, '') else False,
                'name': str(cell(row, 'name') or '').strip() or False,
                'debit': debit,
                'credit': credit,
                'bank_balance': _parse_amount(cell(row, 'balance')) if 'balance' in mapping else 0.0,
            })
        if not vals_list:
            raise UserError('El archivo no tiene movimientos dentro del periodo %s.'
                            % rec._period_label())
        if self.replace_unmatched:
            rec.statement_line_ids.filtered(lambda l: l.match_state == 'unmatched').unlink()
        for seq, vals in enumerate(vals_list, start=1):
            vals['sequence'] = seq
        lines = self.env['construtec.bank.statement.line'].create(vals_list)
        body = 'Estado de cuenta importado (%s): %s movimiento(s).' % (self.filename, len(lines))
        if skipped:
            body += ' %s fuera del periodo, omitidos.' % skipped
        rec.message_post(body=body)
        rec.action_auto_match()
        return {'type': 'ir.actions.act_window_close'}
